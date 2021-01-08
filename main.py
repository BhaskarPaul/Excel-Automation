# # # Created By : Bhaskar # # #

from openpyxl import Workbook
from openpyxl.styles import *
from openpyxl.worksheet.table import Table, TableStyleInfo

# opening txt file
text_file = open("./source.txt")
# forming data from txt file into a list
data = []
# pointer to first of the file
text_file.seek(0)
for i in text_file.readlines():
    data.append(i.rstrip('\n').split(';'))
# creating a new workbook
workbook = Workbook()
# path of excel file where you gonna store data
path = "C:\\Users\\BHASKAR\\Desktop\\Python Automation\\Automate Excel\\DataSheet.xlsx"
workbook.save(path)
# create a variable named sheet
# as default name of sheet is assigned to 'Sheet'
# check the name of sheet using => print(workbook.sheetnames)
sheet = workbook['Sheet']
# change the name of the sheet :)
sheet.title = 'StudentData'
# populating the excel file
for datas in data:
    sheet.append(datas)
# creating a table
table = Table(displayName = 'Table',
              ref = "A1:E6"
)
# generating style
style = TableStyleInfo(name = 'TableStyleMedium9',
                       showRowStripes = True,
                       showColumnStripes = True
)
table.tableStyleInfo = style
# add this table to the sheet
sheet.add_table(table)
# students have roll greater than 12, they will have colored font
font = Font(color = 'a83832',
            bold = True,
            italic = True
)
for row in range(2, 7):
    if int(sheet['D%s' %row].value) > 12:
        sheet['D%s' %row].font = font
# save the workbook and close the text file
workbook.save(path)
text_file.close()
workbook.close()

